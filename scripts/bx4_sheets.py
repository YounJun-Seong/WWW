"""BX4-A2 human sufficiency annotation workbooks (DESIGN_BX4.md). 150 questions per dataset among questions with
|G_inj| = 2 and at least one receiver wrong at (frozen) round 0, drawn with random.Random(20261019). Two blind workbooks
(annotator A / B, different random orders), no model output and no correctness shown; key file kept separately.
  python bx4_sheets.py   -> runs/bx4/human/{annotator_A.xlsx, annotator_B.xlsx, KEY.json}"""
import json
import os
import random

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

from common import ROOT
from rx_analyze import Score, rd
from rx_run import Ctx
import pandas as pd

OUT = f"{ROOT}/runs/bx4/human"
GUIDE = ["Task: decide whether the two passages, ON THEIR OWN, contain enough information to reach the reference answer.",
         "Sufficient: every fact needed to reach the reference answer is stated in the passages (a short inference is fine).",
         "Insufficient: at least one needed fact is missing from the passages.",
         "Ambiguous: the passages only partly support the answer, or the support is unclear.",
         "Do NOT use outside knowledge: an answer you know to be true but cannot read off the passages is Insufficient.",
         "Judge independently; do not discuss items with the other annotator until both sheets are complete."]


def main():
    os.makedirs(OUT, exist_ok=True)
    items = []
    for ds in ("frames", "hotpot"):
        cx = Ctx(ds); g0 = rd(f"runs/final/{ds}_final_gen.jsonl"); c0 = Score(ds).fn(f"runs/final/{ds}_final_judged.jsonl")
        gold = (lambda i: str(pd.read_csv(f"{ROOT}/data/frames_test.tsv", sep="\t").loc[i, "Answer"])) if ds == "frames" else \
               (lambda i, L=json.load(open(f"{ROOT}/data/hotpot_final_labels.json")): L[str(i)])
        elig = sorted(i for i in cx.Q if len(cx.X[i]["G_inj"]) == 2 and (not c0(i, g0[i], "A0") or not c0(i, g0[i], "B0")))
        pick = random.Random(20261019).sample(elig, 150)
        for i in pick:
            P = [cx.pas(x) for x in cx.X[i]["G_inj"]]
            items.append(dict(dataset=ds, idx=i, question=cx.Q[i]["question"], answer=gold(i),
                              p1=f"{P[0]['title']}\n{P[0]['text']}", p2=f"{P[1]['title']}\n{P[1]['text']}"))
        print(f"{ds}: eligible {len(elig)}, sampled 150")
    for k, it in enumerate(items):
        it["item_id"] = f"S{k + 1:03d}"
    json.dump({it["item_id"]: dict(dataset=it["dataset"], idx=it["idx"]) for it in items}, open(f"{OUT}/KEY.json", "w"), indent=1)
    for who, seed in (("A", 1), ("B", 2)):
        order = list(items); random.Random(seed).shuffle(order)
        wb = Workbook(); g = wb.active; g.title = "Guide"
        for r, line in enumerate(GUIDE, 1):
            g.cell(r, 1, line)
        g.column_dimensions["A"].width = 130
        ws = wb.create_sheet("Items")
        hdr = ["item_id", "question", "reference_answer", "passage_1", "passage_2", "label", "note"]
        ws.append(hdr)
        for c in ws[1]:
            c.font = Font(bold=True)
        for it in order:
            ws.append([it["item_id"], it["question"], it["answer"], it["p1"], it["p2"], "", ""])
        dv = DataValidation(type="list", formula1='"Sufficient,Insufficient,Ambiguous"', allow_blank=True)
        ws.add_data_validation(dv); dv.add(f"F2:F{len(order) + 1}")
        for col, w in zip("ABCDEFG", (9, 50, 22, 70, 70, 14, 30)):
            ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        wb.save(f"{OUT}/annotator_{who}.xlsx")
    print(f"saved {len(items)} items x 2 annotators -> {OUT}")


if __name__ == "__main__":
    main()
