"""Excel version of the blind human-validation sheet (same 50 items, same order as data/human_validation_sheet.csv;
the key file is not read). -> data/human_validation_sheet.xlsx"""
import csv

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from common import ROOT

GUIDE = [
    "FRAMES judge validation: 50 items (blind; system and judge verdict hidden)",
    "",
    "Fill only column E 'human_correct' with 1 or 0 (drop-down). Column F 'note' is optional.",
    "Question to answer: does the system response give an answer semantically equivalent to the gold answer?",
    "",
    "1 = same answer in different wording or format ('A horse' vs 'Horse'), or the correct answer plus correct extra detail.",
    "    If the gold cell also contains an explanation, judge only the core answer the question asks for (e.g. item 3: '12').",
    "0 = a different answer; only part of a multi-part answer; no answer ('insufficient information'); several candidates",
    "    without committing to one; or the gold string appears but the stated conclusion differs.",
    "",
    "Judge against the gold answer only; do not look anything up. If you think the gold is wrong, still grade against it",
    "and write 'gold doubtful' in the note. If unsure, still choose 1 or 0 and note why (e.g. 'partial', 'format').",
    "Do not edit other columns or reorder rows. Do not open human_validation_key.json until you are done.",
]


def main():
    rows = list(csv.reader(open(f"{ROOT}/data/human_validation_sheet.csv", newline="")))
    wb = Workbook()
    g = wb.active; g.title = "Guide"
    for i, line in enumerate(GUIDE, 1):
        g.cell(i, 1, line).font = Font(bold=(i == 1), size=12 if i == 1 else 11)
    g.column_dimensions["A"].width = 130
    s = wb.create_sheet("Annotate")
    head = ["item", "question", "gold_answer", "system_response", "human_correct", "note"]
    s.append(head)
    for r in rows[1:]:
        s.append([int(r[0]), r[1], r[2], r[3], None, None])
    for c in s[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
    for col, w in zip("ABCDEF", (6, 60, 30, 60, 15, 30)):
        s.column_dimensions[col].width = w
    for row in s.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        row[4].fill = PatternFill("solid", fgColor="FFF2CC"); row[4].alignment = Alignment(horizontal="center", vertical="top")
    dv = DataValidation(type="list", formula1='"1,0"', allow_blank=True, showErrorMessage=True,
                        errorTitle="1 or 0", error="Enter 1 (equivalent) or 0 (not equivalent).")
    s.add_data_validation(dv); dv.add(f"E2:E{len(rows)}")
    s.freeze_panes = "B2"; s.auto_filter.ref = f"A1:F{len(rows)}"
    wb.active = 1
    wb.save(f"{ROOT}/data/human_validation_sheet.xlsx")
    print(f"xlsx: {len(rows) - 1} items")


if __name__ == "__main__":
    main()
