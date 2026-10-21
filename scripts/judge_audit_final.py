"""Final-cohort judge-sensitivity audit (W5; user request; sampling fixed before any annotation).
From FRAMES final, questions where the judge's verdicts differ between MAS and a comparator are sampled (seed 2027),
with disjoint questions per stratum: 20 MAS–No-comm, 10 MAS–Pooled, 10 MAS–SC. Both policies' responses (exactly the
text the judge saw) become blind items (80), shuffled; the arm and the judge verdict are hidden.
  python judge_audit_final.py  ->  data/judge_audit_final_sheet.xlsx, data/judge_audit_final_key.json"""
import json
import random

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from common import ROOT

STRATA = (("NOCOMM", 20), ("POOL", 10), ("SC", 10))
GUIDE = ["FRAMES judge audit (final cohort): 80 items, blind (system and judge verdict hidden)",
         "", "Same rules as the first validation sheet:",
         "Fill only column E 'human_correct' with 1 or 0. Does the response give an answer semantically equivalent to the gold?",
         "1 = same answer in other wording or format, or the correct answer plus correct extra detail (judge the core answer the question asks for).",
         "0 = different answer; only part of a multi-part answer; no answer; several candidates without committing; or the gold string appears but the conclusion differs.",
         "Grade against the gold only (do not look anything up); write 'gold doubtful' in the note if needed. Do not edit other columns or reorder rows."]


def main():
    df = pd.read_csv(f"{ROOT}/data/frames_test.tsv", sep="\t")
    J = {j["idx"]: j for j in map(json.loads, open(f"{ROOT}/runs/final/frames_final_judged.jsonl"))}
    rng = random.Random(2027); used = set(); items = []
    for other, m in STRATA:
        pool = sorted(i for i, j in J.items() if j["MAS"]["verdict"] != j[other]["verdict"] and i not in used)
        pick = rng.sample(pool, m); used |= set(pick)
        for i in pick:
            for role in ("MAS", other):
                items.append(dict(idx=i, role=role, stratum=f"MAS_vs_{other}", n_discordant_stratum=len(pool) + 0,
                                  judge=bool(J[i][role]["verdict"]), response=J[i][role]["response"]))
    rng.shuffle(items)
    wb = Workbook(); g = wb.active; g.title = "Guide"
    for k, line in enumerate(GUIDE, 1):
        g.cell(k, 1, line).font = Font(bold=(k == 1))
    g.column_dimensions["A"].width = 140
    s = wb.create_sheet("Annotate"); s.append(["item", "question", "gold_answer", "system_response", "human_correct", "note"])
    for n, it in enumerate(items, 1):
        s.append([n, df.loc[it["idx"], "Prompt"], str(df.loc[it["idx"], "Answer"]), it["response"], None, None])
    for c in s[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
    for col, w in zip("ABCDEF", (6, 60, 30, 60, 15, 30)):
        s.column_dimensions[col].width = w
    for row in s.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        row[4].fill = PatternFill("solid", fgColor="FFF2CC")
    dv = DataValidation(type="list", formula1='"1,0"', allow_blank=True); s.add_data_validation(dv); dv.add(f"E2:E{len(items) + 1}")
    s.freeze_panes = "B2"; wb.active = 1
    wb.save(f"{ROOT}/data/judge_audit_final_sheet.xlsx")
    json.dump([dict(item=n, **{k: v for k, v in it.items() if k != "response"}) for n, it in enumerate(items, 1)],
              open(f"{ROOT}/data/judge_audit_final_key.json", "w"), indent=1)
    print(f"audit sheet: {len(items)} items from {len(used)} questions")


if __name__ == "__main__":
    main()
